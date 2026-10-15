"""Super-admin export of the whole user book — one admin's pool or all of it.

Exists because of a real dead end: a super admin's Users page shows their own
direct clients and nothing else, by design. With every client created under a
sub-admin, that page reads "0 users" while the platform carries hundreds —
there was no way to see, let alone hand over, the whole book.

Contacts ARE included here, on the operator's explicit instruction, and only
here: the Users page and every admin API still hand back "••••••••", so a
sub-admin or broker cannot read a client's email or phone anywhere in the
panel. This export is super-admin-only and every download is written to the
audit log with who took which pool — which is what makes the exception
accountable rather than a hole.

A generated placeholder (`<code>@noemail…`, `NOMOB…`) prints as a dash: it is
an absence, not an address, and a column of them reads as data.
"""

from __future__ import annotations

import re
from datetime import timedelta
from io import BytesIO
from typing import Any

from beanie import PydanticObjectId
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.dependencies import _pool_clause, merge_scope
from app.models.position import Position, PositionStatus
from app.models.transaction import TransactionType, WalletTransaction
from app.models.user import User, UserRole, UserStatus
from app.models.wallet import Wallet
from app.services import user_service
from app.utils.time_utils import now_ist, to_utc

# Owner selector values that are not an id.
OWNER_ALL = "all"
OWNER_SELF = "self"

COLUMNS: list[tuple[str, str]] = [
    ("user_code", "Code"),
    ("full_name", "Name"),
    ("email", "Email"),
    ("mobile", "Mobile"),
    ("owner", "Owner"),
    ("status", "Status"),
    ("balance", "Balance"),
    ("credit", "Credit"),
    ("used_margin", "Used margin"),
    ("settlement", "Settlement due"),
    ("weekly_pnl", "Weekly P&L"),
    ("open_positions", "Open positions"),
    ("open_pnl", "Open P&L"),
    ("created_at", "Created"),
    ("last_login_at", "Last login"),
]

_MONEY_KEYS = {"balance", "credit", "used_margin", "settlement", "weekly_pnl", "open_pnl"}


async def _owner_names() -> dict[str, str]:
    rows = await User.find(
        {"role": {"$in": [UserRole.ADMIN.value, UserRole.BROKER.value]}}
    ).to_list()
    return {str(r.id): (r.full_name or r.user_code or str(r.id)) for r in rows}


async def owner_options() -> list[dict[str, str]]:
    """What the picker offers: every admin and broker, plus the two specials."""
    rows = await User.find(
        {"role": {"$in": [UserRole.ADMIN.value, UserRole.BROKER.value]}}
    ).sort("+full_name").to_list()
    out = [
        {"value": OWNER_ALL, "label": "All owners — the whole platform"},
        {"value": OWNER_SELF, "label": "Super admin — direct clients only"},
    ]
    out += [
        {
            "value": str(r.id),
            "label": f"{r.role.value.title()}: {r.full_name or r.user_code}",
        }
        for r in rows
    ]
    return out


async def _query(owner: str, *, include_demo: bool, include_closed: bool) -> dict[str, Any]:
    q: dict[str, Any] = {"role": UserRole.CLIENT.value}
    if not include_closed:
        q["status"] = {"$ne": UserStatus.CLOSED.value}
    if not include_demo:
        q["is_demo"] = {"$ne": True}
        q["email"] = {"$not": re.compile(r"@demo\.local$", re.IGNORECASE)}
    if owner == OWNER_SELF:
        # "Direct clients only", so the broker pin is not redundant: a
        # top-level broker created by the super admin also carries
        # `assigned_admin_id = None`, and so do that broker's clients.
        q["assigned_admin_id"] = None
        q["assigned_broker_id"] = None
    elif owner != OWNER_ALL:
        # This was an `$or` of assigned_admin_id / assigned_broker_id /
        # broker_ancestry against the SAME id, which cannot be right for both
        # roles at once: the last two fields only ever hold BROKER ids. So for
        # an ADMIN owner two of the three arms matched literally nothing, and
        # the sheet quietly dropped every client sitting under a broker that
        # had been transferred INTO that admin's pool — the reassign endpoint
        # mutates only the broker row, so those clients still carry the
        # PREVIOUS admin's `assigned_admin_id`. Ask the shared resolver
        # instead: it is the same clause the Users page and the dashboard
        # counts use, so the export can be reconciled against them rather
        # than disagreeing with both.
        row = await User.get(PydanticObjectId(owner))
        if row is None or row.role not in (UserRole.ADMIN, UserRole.BROKER):
            # `owner_options` offers admins and brokers only. A hand-typed id
            # for anything else exports nothing, rather than falling through
            # to the whole platform under a "Selected owner" title that nobody
            # would read as "all".
            q["_id"] = {"$in": []}
            return q
        # merge_scope, not `q.update`: the admin clause IS an `$or` once that
        # admin has brokers, and a dict holds `$or` only once.
        merge_scope(q, await _pool_clause(row))
    return q


async def collect_rows(
    owner: str = OWNER_ALL,
    *,
    include_demo: bool = False,
    include_closed: bool = False,
) -> list[dict[str, Any]]:
    """One row per user, already formatted for a sheet or a page.

    Everything is batched: wallets, the week's realised P&L and open
    positions each take one round trip for the whole book, not one per user.
    """
    scoped = await _query(owner, include_demo=include_demo, include_closed=include_closed)
    users = await User.find(scoped).sort("-created_at").to_list()
    if not users:
        return []

    ids = [u.id for u in users]
    names = await _owner_names()

    wallets = {
        str(w.user_id): w for w in await Wallet.find({"user_id": {"$in": ids}}).to_list()
    }

    # Current ISO week, Monday 00:00 IST — the same window the users page's
    # "Weekly P&L" column uses, so the two can be reconciled.
    now_i = now_ist()
    week_start = to_utc(
        (now_i - timedelta(days=now_i.weekday())).replace(
            hour=0, minute=0, second=0, microsecond=0
        )
    )
    weekly: dict[str, float] = {}
    try:
        cur = WalletTransaction.get_motor_collection().aggregate(
            [
                {
                    "$match": {
                        "user_id": {"$in": ids},
                        "transaction_type": {
                            "$in": [
                                TransactionType.PNL.value,
                                TransactionType.CHARGES.value,
                                TransactionType.REVERSAL.value,
                                TransactionType.SETTLEMENT_OUTSTANDING_BOOKED.value,
                                TransactionType.SETTLEMENT_OUTSTANDING_RECOVERY.value,
                            ]
                        },
                        "created_at": {"$gte": week_start},
                    }
                },
                {"$group": {"_id": "$user_id", "pnl": {"$sum": {"$toDouble": "$amount"}}}},
            ]
        )
        async for row in cur:
            weekly[str(row["_id"])] = float(row.get("pnl") or 0.0)
    except Exception:  # noqa: BLE001 — a missing column beats a failed export
        weekly = {}

    # Open positions: count and last-known unrealised P&L. Deliberately NOT a
    # live M2M sweep — that is hundreds of feed reads for a file nobody reads
    # tick by tick, and the stored value is what the risk loop last wrote.
    pos_count: dict[str, int] = {}
    pos_pnl: dict[str, float] = {}
    try:
        cur = Position.get_motor_collection().aggregate(
            [
                {"$match": {"user_id": {"$in": ids}, "status": PositionStatus.OPEN.value}},
                {
                    "$group": {
                        "_id": "$user_id",
                        "n": {"$sum": 1},
                        "pnl": {"$sum": {"$toDouble": "$unrealized_pnl"}},
                    }
                },
            ]
        )
        async for row in cur:
            pos_count[str(row["_id"])] = int(row.get("n") or 0)
            pos_pnl[str(row["_id"])] = float(row.get("pnl") or 0.0)
    except Exception:  # noqa: BLE001
        pos_count, pos_pnl = {}, {}

    def _owner_label(u: User) -> str:
        if u.assigned_broker_id:
            return f"Broker: {names.get(str(u.assigned_broker_id), '—')}"
        if u.assigned_admin_id:
            return f"Admin: {names.get(str(u.assigned_admin_id), '—')}"
        return "Super admin"

    def _contact(v: str | None) -> str:
        # Shared formatter, not a local copy of the rule. `reveal=True` is
        # stated rather than assumed: this export is SUPER_ADMIN-only, and
        # the day that gate widens this is the line that has to move.
        return user_service.contact_display(v, reveal=True)

    def _money(v: Any) -> float:
        try:
            return round(float(str(v)), 2)
        except (TypeError, ValueError):
            return 0.0

    rows: list[dict[str, Any]] = []
    for u in users:
        uid = str(u.id)
        w = wallets.get(uid)
        rows.append(
            {
                "user_code": u.user_code or "",
                "full_name": u.full_name or "",
                "email": _contact(u.email),
                "mobile": _contact(u.mobile),
                "owner": _owner_label(u),
                "status": str(u.status.value if hasattr(u.status, "value") else u.status)
                + (" · demo" if u.is_demo else ""),
                "balance": _money(w.available_balance) if w else 0.0,
                "credit": _money(w.credit_limit) if w else 0.0,
                "used_margin": _money(w.used_margin) if w else 0.0,
                "settlement": _money(w.settlement_outstanding) if w else 0.0,
                "weekly_pnl": round(weekly.get(uid, 0.0), 2),
                "open_positions": pos_count.get(uid, 0),
                "open_pnl": round(pos_pnl.get(uid, 0.0), 2),
                "created_at": u.created_at.strftime("%d-%m-%Y %H:%M") if u.created_at else "",
                "last_login_at": (
                    u.last_login_at.strftime("%d-%m-%Y %H:%M") if u.last_login_at else "never"
                ),
            }
        )
    return rows


def _title(owner_label: str, count: int) -> str:
    return f"Users — {owner_label} ({count})"


def to_xlsx(rows: list[dict[str, Any]], owner_label: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"

    head_fill = PatternFill("solid", fgColor="0B1220")
    head_font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=1, value=_title(owner_label, len(rows))).font = Font(
        bold=True, size=13
    )
    for c, (_key, head) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=c, value=head)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(rows, start=3):
        for c, (key, _head) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(key))
            if key in _MONEY_KEYS:
                cell.number_format = "#,##0.00"

    widths = {"full_name": 22, "owner": 24, "email": 14, "created_at": 17, "last_login_at": 17}
    for c, (key, head) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(key, max(11, len(head) + 2))
    ws.freeze_panes = "A3"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rows: list[dict[str, Any]], owner_label: str) -> bytes:
    """Landscape A4, all fifteen columns.

    Email is the one field that can be long enough to break a table, so it —
    with the name and the owner — is wrapped in a Paragraph and given a share
    of the width; everything else is a short number or a code.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    cols = list(COLUMNS)
    # Proportions, not millimetres — scaled to whatever width the page has.
    weights = {
        "user_code": 1.0, "full_name": 1.5, "email": 2.6, "mobile": 1.1,
        "owner": 1.5, "status": 0.9, "balance": 1.1, "credit": 0.9,
        "used_margin": 1.1, "settlement": 1.1, "weekly_pnl": 1.1,
        "open_positions": 0.8, "open_pnl": 1.0, "created_at": 1.3,
        "last_login_at": 1.3,
    }
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=_title(owner_label, len(rows)),
    )
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontSize = 6.5
    body.leading = 8

    data: list[list[Any]] = [[h for _k, h in cols]]
    for row in rows:
        data.append(
            [
                (
                    f"{row.get(k, 0):,.2f}"
                    if k in _MONEY_KEYS
                    else Paragraph(str(row.get(k, "")), body)
                    if k in ("full_name", "owner", "email")
                    else str(row.get(k, ""))
                )
                for k, _h in cols
            ]
        )

    avail = landscape(A4)[0] - 20 * mm
    total_w = sum(weights[k] for k, _h in cols)
    table = Table(
        data,
        colWidths=[avail * weights[k] / total_w for k, _h in cols],
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B1220")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (6, 1), (-3, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )

    story = [
        Paragraph(f"<b>{_title(owner_label, len(rows))}</b>", styles["Title"]),
        Spacer(1, 4 * mm),
        table,
    ]
    doc.build(story)
    return buf.getvalue()
